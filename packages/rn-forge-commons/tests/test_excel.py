"""Tests for rn_forge.commons.excel."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")
pd = pytest.importorskip("pandas")
import rn_forge.commons.excel as excel_module
from openpyxl import Workbook
from openpyxl.styles import Font

from rn_forge.commons.excel import (
    CellFormat,
    ExcelAdapter,
    ExcelUtils,
    WorkbookTemplate,
)

from conftest import raise_


# ---------------------------------------------------------------------------
# CellFormat
# ---------------------------------------------------------------------------


class TestCellFormat:
    def test_default_all_none(self):
        fmt = CellFormat()
        assert fmt.font is None
        assert fmt.fill is None
        assert fmt.border is None
        assert fmt.number_format is None
        assert fmt.alignment is None
        assert fmt.protection is None

    def test_set_number_format(self):
        fmt = CellFormat(number_format="0.00%")
        assert fmt.number_format == "0.00%"

    def test_set_font(self):
        font = Font(bold=True)
        fmt = CellFormat(font=font)
        assert fmt.font is font

    def test_dataclass_mixin_as_dict(self):
        fmt = CellFormat(number_format="0.00%")
        d = fmt.as_dict()
        assert d["number_format"] == "0.00%"

    def test_format_constants_are_cell_formats(self):
        assert isinstance(ExcelUtils.FORMAT_CURRENCY, CellFormat)
        assert isinstance(ExcelUtils.FORMAT_DATE, CellFormat)
        assert isinstance(ExcelUtils.FORMAT_PERCENTAGE, CellFormat)

    def test_format_constants_have_number_format(self):
        assert ExcelUtils.FORMAT_DATE.number_format == "mm/dd/yy;@"
        assert ExcelUtils.FORMAT_PERCENTAGE.number_format == "0.00%"
        assert "$" in ExcelUtils.FORMAT_CURRENCY.number_format


# ---------------------------------------------------------------------------
# WorkbookTemplate
# ---------------------------------------------------------------------------


class TestWorkbookTemplate:
    def test_basic_construction(self):
        tmpl = WorkbookTemplate(name="Sheet1", columns=["A", "B", "C"])
        assert tmpl.name == "Sheet1"
        assert tmpl.columns == ["A", "B", "C"]
        assert tmpl.row_format is None
        assert tmpl.column_formats == {}

    def test_with_formats(self):
        fmt = CellFormat(number_format="0.00%")
        tmpl = WorkbookTemplate(
            name="Data",
            columns=["Name", "Value"],
            row_format=fmt,
            column_formats={2: ExcelUtils.FORMAT_PERCENTAGE},
        )
        assert tmpl.row_format is fmt
        assert tmpl.column_formats[2] is ExcelUtils.FORMAT_PERCENTAGE

    def test_column_formats_default_is_isolated(self):
        t1 = WorkbookTemplate(name="A", columns=[])
        t2 = WorkbookTemplate(name="B", columns=[])
        t1.column_formats[1] = CellFormat()
        assert 1 not in t2.column_formats


# ---------------------------------------------------------------------------
# ExcelUtils — workbook I/O
# ---------------------------------------------------------------------------


class TestWorkbookIO:
    def test_new_workbook(self):
        wb = ExcelUtils.new_workbook()
        assert wb is not None
        assert len(wb.sheetnames) == 1

    def test_new_workbook_write_only(self):
        wb = ExcelUtils.new_workbook(write_only=True)
        assert wb.write_only is True

    def test_write_and_load_file(self, tmp_path: Path):
        wb = Workbook()
        ws = wb.active
        ws.append(["name", "value"])
        ws.append(["alpha", 1])
        path = ExcelUtils.write_workbook(wb, tmp_path / "test.xlsx")
        assert path.exists()
        assert path.suffix == ".xlsx"
        wb2 = ExcelUtils.load_workbook(path)
        assert wb2.active["A1"].value == "name"

    def test_write_workbook_returns_path(self, tmp_path: Path):
        wb = Workbook()
        result = ExcelUtils.write_workbook(wb, tmp_path / "out.xlsx")
        assert isinstance(result, Path)

    def test_write_workbook_bytes(self):
        wb = Workbook()
        buf = ExcelUtils.write_workbook_bytes(wb)
        assert isinstance(buf, BytesIO)
        assert buf.tell() > 0

    def test_load_from_bytes(self, tmp_path: Path):
        wb = Workbook()
        wb.active.append(["x", "y"])
        buf = ExcelUtils.write_workbook_bytes(wb)
        buf.seek(0)
        wb2 = ExcelUtils.load_workbook(buf)
        assert wb2.active["A1"].value == "x"

    def test_write_creates_parent_dirs(self, tmp_path: Path):
        wb = Workbook()
        path = ExcelUtils.write_workbook(wb, tmp_path / "sub" / "dir" / "out.xlsx")
        assert path.exists()

    def test_load_workbook_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        monkeypatch.setattr(
            excel_module.openpyxl,
            "load_workbook",
            lambda *a, **k: raise_(OSError("bad workbook")),
        )
        with pytest.raises(OSError, match="bad workbook"):
            ExcelUtils.load_workbook("bad.xlsx")

    def test_write_workbook_error_propagates(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        wb = Workbook()
        monkeypatch.setattr(wb, "save", lambda *a, **k: raise_(OSError("write failed")))
        with pytest.raises(OSError, match="write failed"):
            ExcelUtils.write_workbook(wb, tmp_path / "out.xlsx")

    def test_write_workbook_bytes_error_propagates(
        self, monkeypatch: pytest.MonkeyPatch
    ):
        wb = Workbook()
        monkeypatch.setattr(
            wb, "save", lambda *a, **k: raise_(OSError("buffer failed"))
        )
        with pytest.raises(OSError, match="buffer failed"):
            ExcelUtils.write_workbook_bytes(wb)


class TestExcelAdapter:
    def test_read_dataframe_returns_sheet_mapping(self, tmp_path: Path):
        source = tmp_path / "book.xlsx"
        with pd.ExcelWriter(source, engine="openpyxl") as writer:
            pd.DataFrame({"name": ["alice"]}).to_excel(
                writer,
                sheet_name="Users",
                index=False,
            )

        workbook = ExcelAdapter.read_dataframe(source)

        assert list(workbook) == ["Users"]
        assert workbook["Users"]["name"].iloc[0] == "alice"

    def test_load_dataframe_reads_from_memory(self) -> None:
        source = BytesIO()
        with pd.ExcelWriter(source, engine="openpyxl") as writer:
            pd.DataFrame({"name": ["alice"]}).to_excel(
                writer,
                sheet_name="Users",
                index=False,
            )
        source.seek(0)

        workbook = ExcelAdapter.load_dataframe(source)

        assert workbook["Users"]["name"].iloc[0] == "alice"

    def test_from_dataframe_returns_workbook(self):
        workbook = ExcelAdapter.from_dataframe(pd.DataFrame({"x": [1, 2, 3]}))

        assert isinstance(workbook, Workbook)
        assert workbook.active["A1"].value == "x"
        assert workbook.active["A2"].value == 1

    def test_from_dataframe_supports_sheet_mapping(self):
        workbook = ExcelAdapter.from_dataframe(
            {
                "Users": pd.DataFrame({"name": ["alice"]}),
                "Roles": pd.DataFrame({"role": ["admin"]}),
            }
        )

        assert workbook["Users"]["A1"].value == "name"
        assert workbook["Roles"]["A2"].value == "admin"

    def test_read_workbook_reads_from_path(self, tmp_path: Path):
        workbook = Workbook()
        workbook.active["A1"] = "name"
        path = ExcelUtils.write_workbook(workbook, tmp_path / "book.xlsx")

        loaded = ExcelAdapter.read_workbook(path)

        assert loaded.active["A1"].value == "name"

    def test_load_workbook_reads_from_memory(self):
        workbook = Workbook()
        workbook.active["A1"] = "name"
        payload = ExcelUtils.write_workbook_bytes(workbook)
        payload.seek(0)

        loaded = ExcelAdapter.load_workbook(payload)

        assert loaded.active["A1"].value == "name"

    def test_write_workbook_writes_to_path(self, tmp_path: Path):
        workbook = Workbook()
        path = ExcelAdapter.write_workbook(workbook, tmp_path / "out.xlsx")

        assert path.exists()

    def test_workbook_to_bytes_returns_buffer(self):
        workbook = Workbook()
        payload = ExcelAdapter.workbook_to_bytes(workbook)

        assert isinstance(payload, BytesIO)
        assert payload.tell() > 0

    def test_read_dataframe_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        monkeypatch.setattr(
            excel_module.pandas,
            "read_excel",
            lambda *a, **k: raise_(RuntimeError("read failed")),
        )
        with pytest.raises(RuntimeError, match="read failed"):
            ExcelAdapter.read_dataframe("book.xlsx")

    def test_load_dataframe_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        monkeypatch.setattr(
            excel_module.pandas,
            "read_excel",
            lambda *a, **k: raise_(RuntimeError("load failed")),
        )
        source = BytesIO()
        with pytest.raises(RuntimeError, match="load failed"):
            ExcelAdapter.load_dataframe(source)

    def test_from_dataframe_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        frame = pd.DataFrame({"x": [1]})

        class FakeWriter:
            def __enter__(self):
                return object()

            def __exit__(self, exc_type, exc, tb):
                return False

        def fail(*args, **kwargs):
            raise RuntimeError("to_excel failed")

        monkeypatch.setattr(
            excel_module.pandas, "ExcelWriter", lambda *a, **k: FakeWriter()
        )
        monkeypatch.setattr(frame, "to_excel", fail)
        with pytest.raises(RuntimeError, match="to_excel failed"):
            ExcelAdapter.from_dataframe(frame)


# ---------------------------------------------------------------------------
# ExcelUtils — format_row
# ---------------------------------------------------------------------------


class TestFormatRow:
    def test_applies_number_format(self):
        wb = Workbook()
        ws = wb.active
        ws.append([1.5, 0.25])
        row = list(ws.iter_rows())[0]
        fmt = CellFormat(number_format="0.00%")
        ExcelUtils.format_row(row, row_format=fmt)
        for cell in row:
            assert cell.number_format == "0.00%"

    def test_column_format_overrides_row_format(self):
        wb = Workbook()
        ws = wb.active
        ws.append([1, 2, 3])
        row = list(ws.iter_rows())[0]
        row_fmt = CellFormat(number_format="0.00")
        col_fmt = CellFormat(number_format="0.00%")
        ExcelUtils.format_row(row, row_format=row_fmt, column_formats={2: col_fmt})
        assert row[0].number_format == "0.00"  # row format
        assert row[1].number_format == "0.00%"  # column override
        assert row[2].number_format == "0.00"  # row format

    def test_none_row_format_leaves_cells_unchanged(self):
        wb = Workbook()
        ws = wb.active
        ws.append([1])
        row = list(ws.iter_rows())[0]
        original = row[0].number_format
        ExcelUtils.format_row(row, row_format=None)
        assert row[0].number_format == original

    def test_none_column_formats_defaults_to_empty(self):
        wb = Workbook()
        ws = wb.active
        ws.append([1])
        row = list(ws.iter_rows())[0]
        # Should not raise
        ExcelUtils.format_row(row, row_format=None, column_formats=None)

    def test_applies_font(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["hello"])
        row = list(ws.iter_rows())[0]
        fmt = CellFormat(font=Font(bold=True))
        ExcelUtils.format_row(row, row_format=fmt)
        assert row[0].font.bold is True


# ---------------------------------------------------------------------------
# ExcelUtils — column_letter
# ---------------------------------------------------------------------------


class TestColumnLetter:
    @pytest.mark.parametrize(
        "number, expected",
        [
            (1, "A"),
            (26, "Z"),
            (27, "AA"),
            (28, "AB"),
            (52, "AZ"),
            (53, "BA"),
            (702, "ZZ"),
            (703, "AAA"),
        ],
    )
    def test_conversion(self, number, expected):
        assert ExcelUtils.column_letter(number) == expected


# ---------------------------------------------------------------------------
# ExcelUtils — table management
# ---------------------------------------------------------------------------


class TestTableManagement:
    def test_add_table(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Value"])
        ws.append(["alpha", 1])
        ExcelUtils.add_table(
            ws, "MyTable", "A1:B2", "TableStyleMedium9", show_row_stripes=True
        )
        assert "MyTable" in ws.tables

    def test_resize_table(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Value"])
        ws.append(["alpha", 1])
        ws.append(["beta", 2])
        ExcelUtils.add_table(ws, "Data", "A1:B2", "TableStyleMedium9")
        ExcelUtils.resize_table(ws, "Data")
        assert ws.tables["Data"].ref == "A1:B3"

    def test_resize_table_expands_to_all_rows(self):
        wb = Workbook()
        ws = wb.active
        for i in range(5):
            ws.append([f"r{i}a", f"r{i}b"])
        ExcelUtils.add_table(ws, "Grow", "A1:B2", "TableStyleMedium9")
        ExcelUtils.resize_table(ws, "Grow")
        assert ws.tables["Grow"].ref == "A1:B5"

    def test_resize_table_with_merged_cell_raises(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name"])
        ws.append(["alpha"])
        ws.append(["beta"])
        ws.merge_cells("A2:A3")
        ExcelUtils.add_table(ws, "Merged", "A1:A2", "TableStyleMedium9")
        with pytest.raises(TypeError, match="MergedCell"):
            ExcelUtils.resize_table(ws, "Merged")


# ---------------------------------------------------------------------------
# ExcelUtils — auto_adjust_column_width
# ---------------------------------------------------------------------------


class TestAutoAdjustColumnWidth:
    def test_sets_column_widths(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["short", "a much longer value here"])
        ExcelUtils.auto_adjust_column_width(ws)
        width_a = ws.column_dimensions["A"].width
        width_b = ws.column_dimensions["B"].width
        assert width_b > width_a

    def test_handles_none_cell_values(self):
        wb = Workbook()
        ws = wb.active
        ws.append([None, "hello", None])
        # Should not raise
        ExcelUtils.auto_adjust_column_width(ws)

    def test_skips_column_when_first_cell_is_merged(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "top"
        ws["B1"] = "header"
        ws["B2"] = "value"
        ws.merge_cells("B1:B2")
        ExcelUtils.auto_adjust_column_width(ws)


# Coverage ROI notes:
# - openpyxl internals (table style serialization, cell object implementations)
#   are exercised via public behavior rather than mocked exhaustively.
