"""Tests for rn_forge.data.pandas."""

from __future__ import annotations

import decimal
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest

pd = pytest.importorskip("pandas")

import rn_forge.data.pandas as pandas_module
from rn_forge.data.pandas import ExcelAdapter, PandasUtils
from rn_forge.data.excel import ExcelUtils
from openpyxl import Workbook
from rn_forge.commons.exceptions import AppException
from rn_forge.commons.testing import raise_


def _row(**kwargs) -> pd.Series:
    """Helper: build a single-row Series from keyword arguments."""
    return pd.DataFrame([kwargs]).iloc[0]


# ---------------------------------------------------------------------------
# get_field
# ---------------------------------------------------------------------------


class TestGetField:
    def test_returns_value(self):
        row = _row(name="alice", age=30)
        assert PandasUtils.get_field(row, "name") == "alice"

    def test_returns_none_for_nan(self):
        row = _row(name=float("nan"))
        result = PandasUtils.get_field(row, "name", required=False)
        assert result is None

    def test_required_nan_raises(self):
        row = _row(name=float("nan"))
        with pytest.raises(AppException, match="name"):
            PandasUtils.get_field(row, "name", required=True)

    def test_required_false_nan_returns_none(self):
        row = _row(name=float("nan"))
        assert PandasUtils.get_field(row, "name", required=False) is None

    def test_required_true_valid_value_does_not_raise(self):
        row = _row(score=99)
        assert PandasUtils.get_field(row, "score") == 99

    def test_numeric_zero_is_valid(self):
        row = _row(count=0)
        assert PandasUtils.get_field(row, "count") == 0

    def test_none_cell_raises_when_required(self):
        row = _row(val=None)
        with pytest.raises(AppException, match="val"):
            PandasUtils.get_field(row, "val", required=True)

    def test_missing_column_raises_key_error(self):
        row = _row(name="alice")
        with pytest.raises(KeyError):
            PandasUtils.get_field(row, "missing")


# ---------------------------------------------------------------------------
# get_string_field
# ---------------------------------------------------------------------------


class TestGetStringField:
    def test_returns_string(self):
        row = _row(city="london")
        assert PandasUtils.get_string_field(row, "city") == "london"

    def test_numeric_coerced_to_string(self):
        row = _row(count=42)
        assert PandasUtils.get_string_field(row, "count") == "42"

    def test_nan_returns_empty_string(self):
        row = _row(name=float("nan"))
        assert PandasUtils.get_string_field(row, "name", required=False) == ""

    def test_required_nan_raises(self):
        row = _row(name=float("nan"))
        with pytest.raises(AppException):
            PandasUtils.get_string_field(row, "name", required=True)


# ---------------------------------------------------------------------------
# get_decimal_field
# ---------------------------------------------------------------------------


class TestGetDecimalField:
    def test_basic_conversion(self):
        row = _row(amount="12.50")
        result = PandasUtils.get_decimal_field(row, "amount")
        assert result == Decimal("12.50")

    def test_default_precision_is_two(self):
        row = _row(amount="12.5678")
        result = PandasUtils.get_decimal_field(row, "amount")
        assert result == Decimal("12.57")

    def test_custom_precision(self):
        row = _row(rate="0.123456")
        result = PandasUtils.get_decimal_field(row, "rate", precision=4)
        assert result == Decimal("0.1235")

    def test_zero_precision(self):
        row = _row(amount="12.9")
        result = PandasUtils.get_decimal_field(row, "amount", precision=0)
        assert result == Decimal("13")

    def test_float_value(self):
        row = _row(amount=1.5)
        result = PandasUtils.get_decimal_field(row, "amount")
        assert result == Decimal("1.50")

    def test_empty_string_returns_zero(self):
        row = _row(amount="")
        result = PandasUtils.get_decimal_field(row, "amount", required=False)
        assert result == Decimal("0.00")

    def test_nan_not_required_returns_zero(self):
        row = _row(amount=float("nan"))
        result = PandasUtils.get_decimal_field(row, "amount", required=False)
        assert result == Decimal("0.00")

    def test_nan_required_raises(self):
        row = _row(amount=float("nan"))
        with pytest.raises(AppException):
            PandasUtils.get_decimal_field(row, "amount", required=True)

    def test_returns_decimal_type(self):
        row = _row(amount="5.00")
        result = PandasUtils.get_decimal_field(row, "amount")
        assert isinstance(result, Decimal)

    def test_invalid_decimal_raises(self):
        row = _row(amount="not-a-number")
        with pytest.raises(decimal.InvalidOperation):
            PandasUtils.get_decimal_field(row, "amount")


class TestIsNa:
    def test_true_for_nan(self):
        assert PandasUtils.is_na(float("nan")) is True

    def test_false_for_regular_value(self):
        assert PandasUtils.is_na("value") is False


# Coverage ROI notes:
# - pandas missing-value semantics are delegated to pandas itself; tests cover
#   integration points and error propagation, not pandas internals.


# ---------------------------------------------------------------------------
# ExcelAdapter
# ---------------------------------------------------------------------------


class TestExcelAdapter:
    def test_read_dataframe_returns_sheet_mapping(self, tmp_path: Path):
        source = tmp_path / "book.xlsx"
        with pd.ExcelWriter(source, engine="openpyxl") as writer:
            pd.DataFrame({"name": ["alice"]}).to_excel(
                writer,
                sheet_name="Users",
                index=False,
            )

        workbook = ExcelAdapter.read_dataframe(source)

        assert list(workbook) == ["Users"]
        assert workbook["Users"]["name"].iloc[0] == "alice"

    def test_load_dataframe_reads_from_memory(self) -> None:
        source = BytesIO()
        with pd.ExcelWriter(source, engine="openpyxl") as writer:
            pd.DataFrame({"name": ["alice"]}).to_excel(
                writer,
                sheet_name="Users",
                index=False,
            )
        source.seek(0)

        workbook = ExcelAdapter.load_dataframe(source)

        assert workbook["Users"]["name"].iloc[0] == "alice"

    def test_from_dataframe_returns_workbook(self):
        workbook = ExcelAdapter.from_dataframe(pd.DataFrame({"x": [1, 2, 3]}))

        assert isinstance(workbook, Workbook)
        assert workbook.active["A1"].value == "x"
        assert workbook.active["A2"].value == 1

    def test_from_dataframe_supports_sheet_mapping(self):
        workbook = ExcelAdapter.from_dataframe(
            {
                "Users": pd.DataFrame({"name": ["alice"]}),
                "Roles": pd.DataFrame({"role": ["admin"]}),
            }
        )

        assert workbook["Users"]["A1"].value == "name"
        assert workbook["Roles"]["A2"].value == "admin"

    def test_read_workbook_reads_from_path(self, tmp_path: Path):
        workbook = Workbook()
        workbook.active["A1"] = "name"
        path = ExcelUtils.write_workbook(workbook, tmp_path / "book.xlsx")

        loaded = ExcelAdapter.read_workbook(path)

        assert loaded.active["A1"].value == "name"

    def test_load_workbook_reads_from_memory(self):
        workbook = Workbook()
        workbook.active["A1"] = "name"
        payload = ExcelUtils.write_workbook_bytes(workbook)
        payload.seek(0)

        loaded = ExcelAdapter.load_workbook(payload)

        assert loaded.active["A1"].value == "name"

    def test_write_workbook_writes_to_path(self, tmp_path: Path):
        workbook = Workbook()
        path = ExcelAdapter.write_workbook(workbook, tmp_path / "out.xlsx")

        assert path.exists()

    def test_workbook_to_bytes_returns_buffer(self):
        workbook = Workbook()
        payload = ExcelAdapter.workbook_to_bytes(workbook)

        assert isinstance(payload, BytesIO)
        assert payload.tell() > 0

    def test_read_dataframe_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        monkeypatch.setattr(
            pd,
            "read_excel",
            lambda *a, **k: raise_(RuntimeError("read failed")),
        )
        with pytest.raises(RuntimeError, match="read failed"):
            ExcelAdapter.read_dataframe("book.xlsx")

    def test_load_dataframe_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        monkeypatch.setattr(
            pd,
            "read_excel",
            lambda *a, **k: raise_(RuntimeError("load failed")),
        )
        source = BytesIO()
        with pytest.raises(RuntimeError, match="load failed"):
            ExcelAdapter.load_dataframe(source)

    def test_from_dataframe_error_propagates(self, monkeypatch: pytest.MonkeyPatch):
        frame = pd.DataFrame({"x": [1]})

        class FakeWriter:
            def __enter__(self):
                return object()

            def __exit__(self, exc_type, exc, tb):
                return False

        def fail(*args, **kwargs):
            raise RuntimeError("to_excel failed")

        monkeypatch.setattr(pd, "ExcelWriter", lambda *a, **k: FakeWriter())
        monkeypatch.setattr(frame, "to_excel", fail)
        with pytest.raises(RuntimeError, match="to_excel failed"):
            ExcelAdapter.from_dataframe(frame)
